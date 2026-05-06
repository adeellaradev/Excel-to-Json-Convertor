import json
import sys
from collections import defaultdict
from datetime import datetime, date, timezone

import openpyxl


# ── Human-readable descriptions for each focus area ─────────────────────────
# "Ax" = Assessment, "Tx" = Treatment, "NAR" = New Assessment Request,
# "NBO" = New Bookings (Online), "DNB" = Did Not Book, "PVA" = Per Visit Average,
# "AHS" = Alberta Health Services (govt-funded), "AR" = Accounts Receivable
FOCUS_META = {
    "Financial": {
        "description": (
            "Revenue billed across all service lines "
            "(PT, RMT, Chiropractic, Pelvic Health). Values in CAD."
        ),
    },
    "Collection": {
        "description": (
            "Payments received, outstanding accounts receivable (AR), "
            "and 4-week collection rate. Values in CAD except rates (0–1 ratio)."
        ),
    },
    "Bookings": {
        "description": "Online appointments booked through the Jane scheduling platform.",
    },
    "Marketing": {
        "description": (
            "New patient volume — total assessments (Ax), consultation requests, "
            "and cumulative Google / Facebook review counts."
        ),
    },
    "Caseload": {
        "description": (
            "Capacity and utilization — available vs booked vs arrived appointments "
            "across PT, RMT, and Chiro. Utilization values are 0–1 ratios."
        ),
    },
    "Treatment Plan": {
        "description": (
            "Treatment plan submission status in Jane — submitted vs unsubmitted plans."
        ),
    },
    "New Clients": {
        "description": "Cancellation counts for new patient assessment appointments.",
    },
    "Online": {
        "description": (
            "Online cancellations and winback attempts for appointments cancelled "
            "through the online portal."
        ),
    },
    "Phone": {
        "description": (
            "Call centre performance via CallHero. "
            "Answer Rate and Booking Rate are 0–1 ratios (e.g. 0.93 = 93%). "
            "NBO = New Bookings (Online), DNB = Did Not Book."
        ),
    },
    "Tx Plan": {
        "description": (
            "Treatment plan prescription and uptake. "
            "NAR = New Assessment Request. "
            "Average Prescribed = avg number of visits prescribed per patient."
        ),
    },
    "Satisfaction": {
        "description": (
            "NPS (Net Promoter Score) collection — how many patients "
            "filled in an NPS survey at their assessment."
        ),
    },
    "Rapport": {
        "description": "Follow-up outcomes — patients re-engaged after dropping off.",
    },
    "Commitment": {
        "description": (
            "Treatment adherence and utilization — PVA (Per Visit Average), "
            "arrived appointments vs available slots across service lines."
        ),
    },
    "AHS": {
        "description": (
            "Alberta Health Services (government-funded) caseload — "
            "AHS assessments and treatments including Hip & Knee class."
        ),
    },
    "Beth": {
        "description": (
            "Beth's outreach metrics — discovery calls, patient database size, "
            "virtual care revenue."
        ),
    },
    "Paula": {
        "description": (
            "Paula's outreach metrics — MD referrals, email campaigns, "
            "open rates, virtual care delivery."
        ),
    },
    "Tracy": {
        "description": (
            "Tracy's outreach metrics — clinician referrals, MD fax campaigns, "
            "discovery bookings."
        ),
    },
    "Tracy+Paula": {
        "description": "Joint Tracy & Paula discovery call delivery.",
    },
    "Uncategorised": {
        "description": (
            "Metrics not yet assigned to a focus area in the source spreadsheet."
        ),
    },
}

# Infer the unit for a metric from its name so consumers know what they're reading
def infer_unit(name: str) -> str:
    n = name.lower()
    if any(w in n for w in ["revenue", " ar", "total ar", "payments", "brace",
                             "orthotics", "pillows", "shockwave revenue",
                             "virtual revenue", "$"]):
        return "CAD"
    if any(w in n for w in [" rate", " %", "utilization", "percentage",
                             "conversion", "open rate"]):
        return "ratio (0–1)"
    return "count"


# Phone sub-groups the spreadsheet splits into 1-metric labels — merge into Phone
MERGE_INTO_PHONE = {"Answer", "Book"}


def normalise_focus(focus: str) -> str:
    return "Phone" if focus in MERGE_INTO_PHONE else focus


def broadcast_merged(ws):
    merged = {}
    for rng in ws.merged_cells.ranges:
        val = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = val
    return merged


def get(ws, row, col, merged):
    return merged.get((row, col), ws.cell(row, col).value)


def clean(val):
    if isinstance(val, str):
        return " ".join(val.split()) or None
    return val


def to_number(val):
    if isinstance(val, str):
        try:
            return float(val)
        except ValueError:
            pass
    return val


def serialise(obj):
    if isinstance(obj, (datetime, date)):
        return obj.isoformat()
    raise TypeError(type(obj))


def convert(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active
    merged = broadcast_merged(ws)

    # ── Internal metric map (includes _col for row reading) ─────────────────
    # Spreadsheet header rows:
    #   1 = category  (sparse — only Phone Performance area)
    #   2 = metric name
    #   3 = focus area
    #   4 = source system
    #   5 = owner (staff name or system responsible for the number)
    raw_metrics = {}
    for col in range(2, ws.max_column + 1):
        name = clean(get(ws, 2, col, merged))
        if not name:
            continue
        raw_focus = clean(get(ws, 3, col, merged)) or "Uncategorised"
        focus     = normalise_focus(raw_focus)
        source    = clean(get(ws, 4, col, merged))
        owner     = clean(get(ws, 5, col, merged))
        raw_metrics[name] = {
            "focus":  focus,
            "source": source,
            "owner":  owner,
            "_col":   col,
        }

    # ── Schema: organised by focus, each section self-describing ────────────
    # Shape:
    #   schema[focus] = {
    #     "description": "...",
    #     "metrics": {
    #       "Metric Name": { "source": "EMR", "unit": "CAD" }   # owner only if set
    #     }
    #   }
    schema: dict = {}
    focus_order = []   # preserve left-to-right column order
    for name, m in raw_metrics.items():
        focus = m["focus"]
        if focus not in schema:
            focus_order.append(focus)
            schema[focus] = {
                "description": FOCUS_META.get(focus, {}).get("description", ""),
                "metrics": {},
            }
        entry: dict = {"source": m["source"], "unit": infer_unit(name)}
        if m["owner"]:
            entry["owner"] = m["owner"]
        schema[focus]["metrics"][name] = entry

    # ── Records: one entry per week, grouped by focus ───────────────────────
    # null means the cell was not filled in that week — it is kept so the
    # structure is consistent and consumers know the field exists.
    # Excel errors (#REF!, #DIV/0!) are normalised to null.
    records = []
    for row in range(1, ws.max_row + 1):
        date_val = ws.cell(row, 1).value
        if not isinstance(date_val, datetime):
            continue

        by_focus: dict = defaultdict(dict)
        for name, m in raw_metrics.items():
            raw = ws.cell(row, m["_col"]).value
            if isinstance(raw, str) and raw.startswith("#"):
                raw = None
            raw = to_number(raw)
            by_focus[m["focus"]][name] = raw

        record: dict = {"date": date_val.strftime("%Y-%m-%d")}
        for focus in focus_order:
            record[focus] = by_focus[focus]

        records.append(record)

    records.sort(key=lambda r: r["date"])

    return {
        "sheet_name":   ws.title,
        "about": (
            "Weekly clinic performance scoreboard. Each record covers one week "
            "and is grouped by focus area. null means the field exists but was not "
            "filled in that week. The full metric catalogue with units and descriptions "
            "lives in 'schema'."
        ),
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "record_count": len(records),
        "metric_count": sum(len(s["metrics"]) for s in schema.values()),
        "schema":       schema,
        "records":      records,
    }


if __name__ == "__main__":
    input_file  = sys.argv[1] if len(sys.argv) > 1 else "Scoreboard Test.xlsx"
    output_file = sys.argv[2] if len(sys.argv) > 2 else "output.json"

    data = convert(input_file)

    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, default=serialise)

    print(
        f"Done: {data['record_count']} records, "
        f"{data['metric_count']} metrics → {output_file}"
    )
